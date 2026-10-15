# -*- coding: utf-8 -*-
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def criar_planilha_kidoo():
    wb = openpyxl.Workbook()
    
    c_header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")  # Indigo escuro
    c_sub_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    c_zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    c_card_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    
    f_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    f_body = Font(name="Segoe UI", size=9, color="0F172A")
    f_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
    f_hook = Font(name="Segoe UI", size=9, italic=True, bold=True, color="4338CA")
    f_green = Font(name="Segoe UI", size=9, bold=True, color="047857")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # ABA 1: LEADS_INBOUND_ENRIQUECIDOS
    ws1 = wb.active
    ws1.title = "LEADS_INBOUND_ENRIQUECIDOS"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:L1")
    ws1["A1"] = "🎯 KIDOO HUB — DIAGNÓSTICO C-LEVEL & VALUE-BASED PRICING"
    ws1["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    ws1["A1"].fill = c_header_fill
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 36

    colunas = [
        "Data/Hora",
        "Contato",
        "Empresa",
        "Nicho",
        "Faturamento / Porte",
        "Nicho Profundo (Modelo de Negócio)",
        "Dor Específica (Gargalo Milionário)",
        "Solução Arquitetada (Kidoo Solution)",
        "Precificação Sugerida (Setup + MRR)",
        "Gancho de Vendas (Pergunta C-Level)",
        "Closer",
        "Status"
    ]

    for col_idx, col_name in enumerate(colunas, 1):
        cell = ws1.cell(row=2, column=col_idx, value=col_name)
        cell.font = f_header
        cell.fill = c_sub_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws1.row_dimensions[2].height = 28

    leads_dados = [
        [
            "13/10/2026 14:30",
            "Carlos Mendonça",
            "Mendonça Locação de Guindastes Pesados",
            "Locação B2B de Equipamentos Industriais",
            "R$ 850.000/mês (Enterprise)",
            "Operação B2B Enterprise de ciclo longo e contratos de alto valor, onde cada ponto de atrito entre marketing e vendas vaza centenas de milhares de reais.",
            "Gargalo milionário de lentidão e falta de governança no atendimento inbound: executivos de conta atendendo contatos desqualificados enquanto contas Enterprise esfriam por demora no primeiro contato.",
            "Kidoo Corporate Lead Routing & RevOps Engine (Triagem Técnica com IA + Algoritmo Round-Robin para Diretores/Closers + Sincronização Bidirecional com CRM corporativo e ERP).",
            "Setup R$ 12.500 e MRR R$ 4.500/mês",
            "\"Carlos, considerando que cada contrato fechado na Mendonça Locação movimenta múltiplos 6 dígitos, quanto dinheiro a sua operação deixa na mesa todo mês por demorar mais de 10 minutos para direcionar um tomador de decisão ao executivo de conta sênior?\"",
            "Pedro",
            "Call Agendada"
        ],
        [
            "13/10/2026 15:10",
            "Dra. Juliana Mendes",
            "Clínica Estética & Dermatologia Savassi",
            "Estética Médica e Harmonização Facial",
            "R$ 180.000/mês (Médio)",
            "Negócio em forte fase de escala e tração, investindo pesado em tráfego pago, mas operando com atendimento centralizado em poucas pessoas no WhatsApp.",
            "Gargalo crítico de conversão: perda de até 30% dos leads nos picos e finais de semana, somado à alta taxa de no-show em agendamentos de procedimentos de alto ticket.",
            "Kidoo Growth Revenue Engine (Agente Conversacional 24/7 com qualificação e quebra de objeções + Sincronização de Agendas + Régua Ativa Anti-No-Show + LP Viva de Alta Conversão).",
            "Setup R$ 5.000 e MRR R$ 1.800/mês",
            "\"Dra. Juliana, com o volume que a clínica já fatura, quando um lead quente de alto ticket manda mensagem no WhatsApp sexta às 21h, quem responde ele em 5 segundos antes que ele feche com o seu principal concorrente?\"",
            "Pedro",
            "Call Agendada"
        ],
        [
            "13/10/2026 15:45",
            "Marcelo 'Barba' Ribeiro",
            "Barbearia Dom Pedro Vintage",
            "Barbearia Premium e Cuidados Masculinos",
            "R$ 35.000/mês (Micro)",
            "Operação local enxuta onde os donos ou recepcionistas acumulam atendimento ao cliente, execução operacional e cobranças.",
            "Tempo excessivo gasto digitando mensagens manuais no WhatsApp, perda de clientes fora do expediente comercial e estagnação da base antiga sem recompra.",
            "Kidoo Smart Scheduling & WhatsApp AI (Agente Conversacional 24/7 + Sincronização com Google Calendar + Confirmação Automática de Horários).",
            "Setup R$ 2.500 e MRR R$ 750/mês",
            "\"Marcelo, hoje quantos clientes deixam de agendar na Dom Pedro simplesmente porque mandam mensagem à noite ou no final de semana e não encontram ninguém para responder na hora?\"",
            "Arthur",
            "Novo Lead"
        ]
    ]

    for row_idx, lead in enumerate(leads_dados, 3):
        ws1.row_dimensions[row_idx].height = 56
        is_zebra = (row_idx % 2 == 0)
        for col_idx, valor in enumerate(lead, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=valor)
            cell.font = f_body
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col_idx in [1, 5, 11, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if col_idx == 8:
                cell.font = f_bold
            if col_idx == 9:
                cell.font = f_green
            if col_idx == 10:
                cell.font = f_hook
            if is_zebra:
                cell.fill = c_zebra_fill

    col_widths = {
        "A": 16, "B": 20, "C": 30, "D": 26, "E": 20,
        "F": 34, "G": 36, "H": 40, "I": 26, "J": 44,
        "K": 14, "L": 16
    }
    for col_letter, width in col_widths.items():
        ws1.column_dimensions[col_letter].width = width

    # ABA 2: REGRAS_VALUE_BASED_PRICING
    ws2 = wb.create_sheet(title="VALUE_BASED_PRICING")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:E1")
    ws2["A1"] = "💰 KIDOO HUB — MATRIZ DE PRECIFICAÇÃO BASEADA EM VALOR"
    ws2["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    ws2["A1"].fill = c_header_fill
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 36

    colunas_pricing = ["Faixa de Faturamento", "Perfil do Lead", "Setup Recomendado", "MRR Recomendado", "Diretriz Comercial C-Level"]
    for col_idx, col_name in enumerate(colunas_pricing, 1):
        cell = ws2.cell(row=2, column=col_idx, value=col_name)
        cell.font = f_header
        cell.fill = c_sub_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws2.row_dimensions[2].height = 26

    regras = [
        ["Micro (Até R$ 50k/mês)", "Negócios locais, barbearias, consultórios enxutos", "R$ 2.000 a R$ 3.000", "R$ 600 a R$ 800/mês", "Venda alívio operacional e captura de clientes fora de hora."],
        ["Médio (R$ 100k a R$ 300k/mês)", "Clínicas estruturadas, e-commerces, distribuidoras", "R$ 4.000 a R$ 6.000", "R$ 1.500 a R$ 2.500/mês", "Venda aceleração de receita, lead scoring e esteira de vendas."],
        ["Enterprise / Corporate (+R$ 500k/mês)", "Grandes locadoras B2B, indústrias, construtoras, redes", "R$ 10.000 a R$ 15.000", "R$ 3.500 a R$ 6.000/mês", "Venda governança, exclusividade e impacto direto na última linha do balanço."]
    ]

    for row_idx, item in enumerate(regras, 3):
        ws2.row_dimensions[row_idx].height = 32
        is_zebra = (row_idx % 2 == 0)
        for col_idx, valor in enumerate(item, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=valor)
            cell.font = f_body
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if col_idx in [1, 3, 4]:
                cell.font = f_bold
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if is_zebra:
                cell.fill = c_zebra_fill

    col_widths_ws2 = {"A": 26, "B": 38, "C": 22, "D": 22, "E": 48}
    for col_letter, width in col_widths_ws2.items():
        ws2.column_dimensions[col_letter].width = width

    # ABA 3: DASHBOARD_REVOPS
    ws3 = wb.create_sheet(title="DASHBOARD_REVOPS")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:E1")
    ws3["A1"] = "📊 KIDOO HUB — PAINEL EXECUTIVO DE REVOPS & PIPELINE"
    ws3["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    ws3["A1"].fill = c_header_fill
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 36

    ws3["B3"] = "TOTAL DE LEADS DIAGNOSTICADOS"
    ws3["B3"].font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
    ws3["B4"] = "=COUNTA(LEADS_INBOUND_ENRIQUECIDOS!B3:B100)"
    ws3["B4"].font = Font(name="Segoe UI", size=20, bold=True, color="1E1B4B")
    ws3["B4"].fill = c_card_fill
    ws3["B4"].alignment = Alignment(horizontal="center", vertical="center")

    ws3["D3"] = "VALOR TOTAL DE SETUP NO PIPELINE"
    ws3["D3"].font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
    ws3["D4"] = "R$ 20.000,00"
    ws3["D4"].font = Font(name="Segoe UI", size=20, bold=True, color="047857")
    ws3["D4"].fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    ws3["D4"].alignment = Alignment(horizontal="center", vertical="center")

    ws3.row_dimensions[3].height = 18
    ws3.row_dimensions[4].height = 36

    caminho_saida = r"C:\Users\pedro\.gemini\antigravity-ide\scratch\kidoo-antigravity-engine\PLANILHA_KIDOO_REVOPS_INBOUND.xlsx"
    wb.save(caminho_saida)
    print(f"[SUCESSO] Planilha C-Level gerada em: {caminho_saida}")
    return caminho_saida

if __name__ == "__main__":
    criar_planilha_kidoo()
